from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).parent
RULES_PATH = ROOT / "rules" / "automobile_loan_rules.json"
MAX_FILE_BYTES = 15 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def load_rules() -> dict[str, Any]:
    with RULES_PATH.open("r", encoding="utf-8") as f:
        return json.load(f)


def extract_text(filename: str, raw: bytes) -> tuple[str, str | None]:
    """Extract embedded text locally. This does not perform OCR."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".txt":
            return raw.decode("utf-8", errors="replace"), None
        if suffix == ".docx":
            doc = Document(io.BytesIO(raw))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts), None
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            parts = [(page.extract_text() or "") for page in reader.pages]
            text = "\n".join(parts)
            warning = None if text.strip() else "No embedded text found. This may be a scanned PDF; OCR is not included."
            return text, warning
    except Exception as exc:
        return "", f"Could not extract text: {type(exc).__name__}: {exc}"
    return "", "Unsupported file type."


def classify_document(filename: str, text: str, rules: dict[str, Any]) -> str:
    haystack = f"{filename} {text[:4000]}".lower()
    scores: dict[str, int] = {}
    for doc_type, config in rules["document_types"].items():
        scores[doc_type] = sum(1 for keyword in config["keywords"] if keyword.lower() in haystack)
    best = max(scores, key=scores.get) if scores else "other"
    return best if scores.get(best, 0) > 0 else "other"


def extract_first(patterns: list[str], text: str, flags: int = re.IGNORECASE) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text, flags)
        if match:
            return re.sub(r"\s+", " ", match.group(1)).strip(" \t:,-")
    return None


def extract_fields(text: str) -> dict[str, str | None]:
    name = extract_first([
        r"(?:borrower|applicant|customer)\s*(?:name)?\s*[:\-]\s*([A-Za-z][A-Za-z .'-]{2,80})",
        r"(?:name of borrower|name of applicant)\s*[:\-]\s*([A-Za-z][A-Za-z .'-]{2,80})",
    ], text)
    amount = extract_first([
        r"(?:loan amount|facility amount|proposed limit|loan limit)\s*[:\-]?\s*(?:NPR|Rs\.?|रु\.?|रू\.?)?\s*([0-9][0-9,]*(?:\.\d{1,2})?)",
    ], text)
    return {"borrower_name": name, "facility_amount": amount}


def add_finding(findings: list[dict[str, str]], rule_id: str, severity: str,
                status: str, title: str, detail: str, next_step: str) -> None:
    findings.append({
        "rule_id": rule_id,
        "severity": severity,
        "status": status,
        "title": title,
        "detail": detail,
        "recommended_action": next_step,
    })


def review_documents(documents: list[dict[str, Any]], rules: dict[str, Any]) -> dict[str, Any]:
    findings: list[dict[str, str]] = []
    if not documents:
        add_finding(findings, "DOC-000", "high", "review_required", "No documents uploaded",
                    "The review was run without any documents.", "Upload the available credit-file documents.")
        return {"findings": findings, "summary": {"documents": 0, "high": 1, "review_required": 1, "information": 0}}

    present_types = {d["document_type"] for d in documents}
    for required in rules["required_document_types"]:
        if required not in present_types:
            label = rules["document_types"].get(required, {}).get("label", required)
            add_finding(findings, f"DOC-{required.upper()}", "medium", "review_required",
                        f"Document not detected: {label}",
                        "No uploaded document was classified as this type. Classification can be wrong; verify manually.",
                        f"Confirm whether {label} is required for this product and obtain/check it if applicable.")

    all_text = "\n".join(d["text"] for d in documents)
    all_lower = all_text.lower()
    for rule in rules["text_checks"]:
        found = any(term.lower() in all_lower for term in rule["search_terms"])
        if rule["expected"] == "present" and not found:
            add_finding(findings, rule["id"], rule["severity"], "review_required", rule["title"],
                        "The expected term/evidence was not detected by simple text search. It may be worded differently or present in a scanned page.",
                        rule["action"])
        elif rule["expected"] == "absent" and found:
            add_finding(findings, rule["id"], rule["severity"], "review_required", rule["title"],
                        "A term associated with this control was detected. This is only a screening flag, not a determination.",
                        rule["action"])

    extracted = []
    for d in documents:
        fields = extract_fields(d["text"])
        extracted.append({"filename": d["filename"], "document_type": d["document_type"], **fields})

    names = [(x["filename"], x["borrower_name"]) for x in extracted if x["borrower_name"]]
    normalized_names = {re.sub(r"[^a-z]", "", name.lower()) for _, name in names}
    if len(normalized_names) > 1:
        add_finding(findings, "XDOC-NAME", "high", "review_required", "Possible borrower-name mismatch",
                    "Extracted borrower/applicant names differ across documents: " +
                    "; ".join(f"{fn}: {name}" for fn, name in names),
                    "Compare the original documents and confirm the correct legal name.")

    amounts = []
    for item in extracted:
        if item["facility_amount"]:
            try:
                amounts.append((item["filename"], float(item["facility_amount"].replace(",", ""))))
            except ValueError:
                pass
    distinct_amounts = {amount for _, amount in amounts}
    if len(distinct_amounts) > 1:
        add_finding(findings, "XDOC-AMOUNT", "high", "review_required", "Possible facility-amount mismatch",
                    "Different amounts were extracted: " +
                    "; ".join(f"{fn}: {amount:,.2f}" for fn, amount in amounts),
                    "Compare the sanction/proposal and supporting documents; confirm currency, amount, and facility context.")

    for d in documents:
        if len(d["text"].strip()) < 40:
            add_finding(findings, "TXT-LOW-" + re.sub(r"[^A-Z0-9]", "", d["filename"].upper())[:20],
                        "medium", "review_required", f"Low text extracted: {d['filename']}",
                        "The extracted text is very short. The file may be scanned, image-only, encrypted, or incorrectly formatted.",
                        "Open the original document and check readability; use bank-approved OCR if appropriate.")

    if not findings:
        add_finding(findings, "INFO-001", "info", "information", "No prototype flags generated",
                    "The automated screening did not identify a configured flag. This does not mean the file is compliant or complete.",
                    "CAD must still review the source documents and approved product policy.")

    summary = {
        "documents": len(documents),
        "high": sum(1 for f in findings if f["severity"] == "high"),
        "review_required": sum(1 for f in findings if f["status"] == "review_required"),
        "information": sum(1 for f in findings if f["severity"] == "info"),
    }
    return {"findings": findings, "extracted_fields": extracted, "summary": summary}


st.set_page_config(page_title="Credit File Review MVP", page_icon="📄", layout="wide")
st.title("AI-Assisted Credit File Review")
st.caption("Automobile-loan MVP • Local text extraction + configurable rules • Human review required")

st.warning(
    "Prototype only. Use synthetic or approved anonymized test documents. Do not upload real customer "
    "credit files unless your bank has approved this environment. This app does not approve loans, "
    "replace CAD judgment, or validate compliance with current bank policy."
)

rules = load_rules()
with st.sidebar:
    st.header("Review settings")
    st.write("Product:", rules["product_name"])
    st.write("Rules version:", rules["rules_version"])
    st.write("Processing:", "Local in this app process")
    st.write("Supported:", "PDF, DOCX, TXT")
    st.write("Max per file:", "15 MB")
    st.info("Scanned PDF OCR and external LLM calls are not included in this MVP.")

uploads = st.file_uploader(
    "Upload credit-file documents",
    type=["pdf", "docx", "txt"],
    accept_multiple_files=True,
    help="Upload only synthetic or bank-approved anonymized test files.",
)

if st.button("Review credit file", type="primary", disabled=not uploads):
    docs: list[dict[str, Any]] = []
    extraction_warnings: list[str] = []
    for upload in uploads or []:
        if Path(upload.name).suffix.lower() not in SUPPORTED_EXTENSIONS:
            extraction_warnings.append(f"{upload.name}: unsupported file type.")
            continue
        raw = upload.getvalue()
        if len(raw) > MAX_FILE_BYTES:
            extraction_warnings.append(f"{upload.name}: file exceeds the 15 MB prototype limit.")
            continue
        text, warning = extract_text(upload.name, raw)
        doc_type = classify_document(upload.name, text, rules)
        docs.append({
            "filename": upload.name,
            "document_type": doc_type,
            "text": text,
            "bytes": len(raw),
        })
        if warning:
            extraction_warnings.append(f"{upload.name}: {warning}")

    report = review_documents(docs, rules)
    report["metadata"] = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "product": rules["product_name"],
        "rules_version": rules["rules_version"],
        "prototype": True,
        "human_review_required": True,
    }
    report["documents"] = [
        {"filename": d["filename"], "detected_type": d["document_type"],
         "bytes": d["bytes"], "characters_extracted": len(d["text"])}
        for d in docs
    ]
    st.session_state["review_report"] = report
    st.session_state["extraction_warnings"] = extraction_warnings

if st.session_state.get("extraction_warnings"):
    st.subheader("Extraction notes")
    for warning in st.session_state["extraction_warnings"]:
        st.warning(warning)

report = st.session_state.get("review_report")
if report:
    st.subheader("Review summary")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Documents processed", report["summary"]["documents"])
    c2.metric("High-priority flags", report["summary"]["high"])
    c3.metric("Review-required findings", report["summary"]["review_required"])
    c4.metric("Information flags", report["summary"]["information"])

    st.subheader("Document inventory")
    if report.get("documents"):
        st.dataframe(report["documents"], use_container_width=True, hide_index=True)

    st.subheader("Findings")
    for finding in report["findings"]:
        title = f'{finding["severity"].upper()} · {finding["title"]}'
        with st.expander(title, expanded=finding["severity"] == "high"):
            st.write("Status:", finding["status"])
            st.write("Rule:", finding["rule_id"])
            st.write(finding["detail"])
            st.write("Suggested next step:", finding["recommended_action"])

    with st.expander("Extracted fields (verify against originals)"):
        st.dataframe(report.get("extracted_fields", []), use_container_width=True, hide_index=True)

    st.download_button(
        "Download review report (JSON)",
        data=json.dumps(report, ensure_ascii=False, indent=2),
        file_name="credit_file_review_report.json",
        mime="application/json",
    )

st.divider()
st.caption("Every flag is a prompt for verification. No flags is not a credit approval or a compliance certificate.")
